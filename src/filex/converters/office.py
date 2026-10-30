# # """Local DOCX and PDF conversion."""

# # import shutil
# # import subprocess
# # import tempfile
# # from pathlib import Path

# # import pymupdf
# # from docx import Document
# # from docx.enum.text import WD_BREAK
# # from docx.shared import Inches
# # from pdf2docx import Converter

# # from filex.exceptions import FileXError

# # def find_libreoffice() -> Path:
# #     """Find the LibreOffice executable."""
# #     executable = shutil.which("soffice")

# #     if executable:
# #         return Path(executable)

# #     candidates = [
# #         Path("C:/Program Files/LibreOffice/program/soffice.exe"),
# #         Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
# #     ]

# #     for candidate in candidates:
# #         if candidate.is_file():
# #             return candidate

# #     raise FileXError(
# #         "LibreOffice was not found. It is required for DOCX-to-PDF conversion."
# #     )


# # class OfficeConverter:
# #     """Convert PDF to DOCX or DOCX to PDF."""

# #     def convert(
# #         self,
# #         input_path: Path,
# #         output_path: Path,
# #         target_format: str,
# #         mode: str = "editable",
# #     ) -> None:  
# #         input_path = input_path.resolve()
# #         output_path = output_path.resolve()
# #         source_format = input_path.suffix.lower().lstrip(".")

# #         # if source_format == "pdf" and target_format == "docx":
# #         #     self._pdf_to_docx(input_path, output_path)
# #         #     return

# #         if source_format == "pdf" and target_format == "docx":
# #             if mode == "visual":
# #                 self._pdf_to_docx_visual(input_path, output_path)
# #             elif mode == "editable":
# #                 self._pdf_to_docx(input_path, output_path)
# #             else:
# #                 raise FileXError(
# #                     f"Unsupported PDF-to-DOCX mode: {mode}."
# #                 )
# #             return

# #         if source_format == "docx" and target_format == "pdf":
# #             self._docx_to_pdf(input_path, output_path)
# #             return

# #         raise FileXError(
# #             f"Unsupported document conversion: "
# #             f"{source_format} to {target_format}."
# #         )

# #     def _pdf_to_docx(
# #         self,
# #         input_path: Path,
# #         output_path: Path,
# #     ) -> None:
# #         """Convert PDF to DOCX using pdf2docx."""
# #         with tempfile.TemporaryDirectory(prefix="filex_") as temp:
# #             generated_file = Path(temp) / "converted.docx"
# #             converter = None

# #             try:
# #                 converter = Converter(str(input_path))
# #                 converter.convert(str(generated_file))
# #             except Exception as exc:
# #                 raise FileXError(
# #                     f"PDF-to-DOCX conversion failed: {exc}"
# #                 ) from exc
# #             finally:
# #                 if converter is not None:
# #                     converter.close()

# #             if not generated_file.is_file():
# #                 raise FileXError(
# #                     "PDF-to-DOCX conversion did not produce an output file."
# #                 )

# #             self._save_without_overwrite(generated_file, output_path)

# #     def _docx_to_pdf(
# #         self,
# #         input_path: Path,
# #         output_path: Path,
# #     ) -> None:
# #         """Convert DOCX to PDF using LibreOffice."""
# #         soffice = find_libreoffice()

# #         with tempfile.TemporaryDirectory(prefix="filex_") as temp:
# #             temp_dir = Path(temp)
# #             result_dir = temp_dir / "result"
# #             profile_dir = temp_dir / "profile"
# #             result_dir.mkdir()
# #             profile_dir.mkdir()

# #             command = [
# #                 str(soffice),
# #                 "--headless",
# #                 f"-env:UserInstallation={profile_dir.as_uri()}",
# #                 "--convert-to",
# #                 "pdf",
# #                 "--outdir",
# #                 str(result_dir),
# #                 str(input_path),
# #             ]

# #             try:
# #                 process = subprocess.run(
# #                     command,
# #                     capture_output=True,
# #                     text=True,
# #                     timeout=120,
# #                     check=False,
# #                 )
# #             except subprocess.TimeoutExpired as exc:
# #                 raise FileXError(
# #                     "Document conversion timed out after 120 seconds."
# #                 ) from exc
# #             except OSError as exc:
# #                 raise FileXError(
# #                     f"Could not start LibreOffice: {exc}"
# #                 ) from exc

# #             generated_file = result_dir / f"{input_path.stem}.pdf"

# #             if process.returncode != 0 or not generated_file.is_file():
# #                 details = (process.stderr or process.stdout).strip()
# #                 raise FileXError(
# #                     "DOCX-to-PDF conversion failed."
# #                     + (f" Details: {details}" if details else "")
# #                 )

# #             self._save_without_overwrite(generated_file, output_path)

# #     def _pdf_to_docx_visual(
# #         self,
# #         input_path: Path,
# #         output_path: Path,
# #     ) -> None:
# #         """Convert PDF pages into images embedded in a DOCX file."""
# #         input_path = input_path.resolve()
# #         output_path = output_path.resolve()

# #         try:
# #             with pymupdf.open(input_path) as pdf:
# #                 if pdf.page_count == 0:
# #                     raise FileXError(
# #                         "The input PDF contains no pages."
# #                     )

# #                 document = Document()
# #                 section = document.sections[0]

# #                 # Use the PDF page dimensions and remove Word margins.
# #                 first_page = pdf[0]
# #                 page_width = Inches(first_page.rect.width / 72)
# #                 page_height = Inches(first_page.rect.height / 72)

# #                 section.page_width = page_width
# #                 section.page_height = page_height
# #                 section.top_margin = Inches(0)
# #                 section.bottom_margin = Inches(0)
# #                 section.left_margin = Inches(0)
# #                 section.right_margin = Inches(0)

# #                 with tempfile.TemporaryDirectory(
# #                     prefix="filex_visual_"
# #                 ) as temp:
# #                     temp_dir = Path(temp)

# #                     for index, page in enumerate(pdf):
# #                         # Render the PDF page at 2x resolution.
# #                         pixmap = page.get_pixmap(
# #                             matrix=pymupdf.Matrix(2, 2),
# #                             alpha=False,
# #                         )

# #                         image_path = (
# #                             temp_dir / f"page_{index + 1}.png"
# #                         )
# #                         pixmap.save(str(image_path))

# #                         # Put one page image into the Word document.
# #                         # paragraph = (
# #                         #     document.paragraphs[0]
# #                         #     if index == 0
# #                         #     else document.add_paragraph()
# #                         # )

# #                         paragraph = document.add_paragraph()
                        
# #                         paragraph.paragraph_format.space_before = 0
# #                         paragraph.paragraph_format.space_after = 0

# #                         run = paragraph.add_run()
# #                         run.add_picture(
# #                             str(image_path),
# #                             width=page_width,
# #                             height=page_height,
# #                         )

# #                         # Start the next PDF page on a new Word page.
# #                         if index < pdf.page_count - 1:
# #                             run.add_break(WD_BREAK.PAGE)

# #                     # Save temporarily before using FileX's save method.
# #                     temp_output = temp_dir / "visual_output.docx"
# #                     document.save(str(temp_output))

# #                     self._save_without_overwrite(
# #                         temp_output,
# #                         output_path,
# #                     )

# #         except FileXError:
# #             raise
# #         except Exception as exc:
# #             raise FileXError(
# #                 "PDF-to-DOCX visual conversion failed: "
# #                 f"{exc}"
# #             ) from exc

    
# #     @staticmethod
# #     def _save_without_overwrite(
# #         source_path: Path,
# #         output_path: Path,
# #     ) -> None:
# #         """Copy a converted file without overwriting an existing file."""
# #         output_created = False

# #         try:
# #             with source_path.open("rb") as source:
# #                 with output_path.open("xb") as destination:
# #                     output_created = True
# #                     shutil.copyfileobj(source, destination)
# #         except FileExistsError as exc:
# #             raise FileXError(
# #                 f"Output file already exists: {output_path}. "
# #                 "Choose another output path."
# #             ) from exc
# #         except OSError as exc:
# #             if output_created:
# #                 output_path.unlink(missing_ok=True)

# #             raise FileXError(
# #                 f"Could not save converted document: {exc}"
# #             ) from exc





# """Document conversion powered by LibreOffice, with a specialized PDF importer."""

# import shutil
# import subprocess
# import tempfile
# from pathlib import Path

# from filex.exceptions import FileXError


# def find_libreoffice() -> Path:
#     """Find the LibreOffice executable."""
#     executable = shutil.which("soffice")
#     if executable:
#         return Path(executable)

#     candidates = [
#         Path("C:/Program Files/LibreOffice/program/soffice.exe"),
#         Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
#     ]
#     for candidate in candidates:
#         if candidate.is_file():
#             return candidate

#     raise FileXError(
#         "LibreOffice was not found. Install LibreOffice and make sure "
#         "soffice is available in PATH, or install it in a standard location."
#     )


# class OfficeConverter:
#     """Convert supported document formats using local conversion engines."""

#     def convert(
#         self,
#         input_path: Path,
#         output_path: Path,
#         target_format: str,
#         mode: str = "editable",
#     ) -> None:
#         input_path = input_path.resolve()
#         output_path = output_path.resolve()
#         source_format = input_path.suffix.lower().lstrip(".")

#         if source_format == "pdf" and target_format == "docx":
#             if mode == "visual":
#                 self._pdf_to_docx_visual(input_path, output_path)
#             elif mode == "editable":
#                 self._pdf_to_docx(input_path, output_path)
#             else:
#                 raise FileXError(f"Unsupported PDF-to-DOCX mode: {mode}")
#             return

#         self._convert_with_libreoffice(
#             input_path=input_path,
#             output_path=output_path,
#             target_format=target_format,
#         )

#     def _convert_with_libreoffice(
#         self,
#         input_path: Path,
#         output_path: Path,
#         target_format: str,
#     ) -> None:
#         """Convert a document to any target format supported by LibreOffice."""
#         soffice = find_libreoffice()

#         with tempfile.TemporaryDirectory(prefix="filex_") as temp:
#             temp_dir = Path(temp)
#             result_dir = temp_dir / "result"
#             profile_dir = temp_dir / "profile"
#             result_dir.mkdir()
#             profile_dir.mkdir()

#             command = [
#                 str(soffice),
#                 "--headless",
#                 f"-env:UserInstallation={profile_dir.as_uri()}",
#                 "--convert-to",
#                 target_format,
#                 "--outdir",
#                 str(result_dir),
#                 str(input_path),
#             ]

#             try:
#                 process = subprocess.run(
#                     command,
#                     capture_output=True,
#                     text=True,
#                     timeout=120,
#                     check=False,
#                 )
#             except subprocess.TimeoutExpired as exc:
#                 raise FileXError(
#                     "Document conversion timed out after 120 seconds."
#                 ) from exc
#             except OSError as exc:
#                 raise FileXError(
#                     f"Could not start LibreOffice: {exc}"
#                 ) from exc

#             generated_file = result_dir / f"{input_path.stem}.{target_format}"
#             if process.returncode != 0 or not generated_file.is_file():
#                 details = (process.stderr or process.stdout).strip()
#                 raise FileXError(
#                     f"Conversion from {input_path.suffix.lstrip('.').upper()} "
#                     f"to {target_format.upper()} failed. LibreOffice may not "
#                     "support this format pair."
#                     + (f" Details: {details}" if details else "")
#                 )

#             self._save_without_overwrite(generated_file, output_path)

#     def _pdf_to_docx(
#         self,
#         input_path: Path,
#         output_path: Path,
#     ) -> None:
#         """Convert PDF to editable DOCX using pdf2docx."""
#         from pdf2docx import Converter

#         with tempfile.TemporaryDirectory(prefix="filex_") as temp:
#             generated_file = Path(temp) / "converted.docx"
#             converter = None
#             try:
#                 converter = Converter(str(input_path))
#                 converter.convert(str(generated_file))
#             except Exception as exc:
#                 raise FileXError(
#                     f"PDF-to-DOCX conversion failed: {exc}"
#                 ) from exc
#             finally:
#                 if converter is not None:
#                     converter.close()

#             if not generated_file.is_file():
#                 raise FileXError(
#                     "PDF-to-DOCX conversion did not produce an output file."
#                 )
#             self._save_without_overwrite(generated_file, output_path)

#     def _pdf_to_docx_visual(
#         self,
#         input_path: Path,
#         output_path: Path,
#     ) -> None:
#         """Render PDF pages as images inside a DOCX document."""
#         try:
#             import pymupdf
#             from docx import Document
#             from docx.shared import Inches

#             with tempfile.TemporaryDirectory(prefix="filex_visual_") as temp:
#                 temp_dir = Path(temp)
#                 generated_file = temp_dir / "converted.docx"
#                 document = Document()

#                 with pymupdf.open(input_path) as pdf_document:
#                     if pdf_document.page_count == 0:
#                         raise FileXError("The input PDF contains no pages.")

#                     for page in pdf_document:
#                         pixmap = page.get_pixmap(
#                             matrix=pymupdf.Matrix(2, 2),
#                             alpha=False,
#                         )
#                         image_path = temp_dir / f"page_{page.number}.png"
#                         pixmap.save(image_path)
#                         document.add_picture(
#                             str(image_path),
#                             width=Inches(6.5),
#                         )

#                 document.save(generated_file)
#                 self._save_without_overwrite(generated_file, output_path)
#         except FileXError:
#             raise
#         except Exception as exc:
#             raise FileXError(
#                 f"PDF-to-DOCX visual conversion failed: {exc}"
#             ) from exc

#     @staticmethod
#     def _save_without_overwrite(
#         source_path: Path,
#         output_path: Path,
#     ) -> None:
#         """Copy a converted file without overwriting an existing file."""
#         output_created = False
#         try:
#             with source_path.open("rb") as source:
#                 with output_path.open("xb") as destination:
#                     output_created = True
#                     shutil.copyfileobj(source, destination)
#         except FileExistsError as exc:
#             raise FileXError(
#                 f"Output file already exists: {output_path}. "
#                 "Choose another output path."
#             ) from exc
#         except OSError as exc:
#             if output_created:
#                 output_path.unlink(missing_ok=True)
#             raise FileXError(
#                 f"Could not save converted document: {exc}"
#             ) from exc




"""Document conversion powered by LibreOffice, with a specialized PDF importer."""

import shutil
import subprocess
import tempfile
from pathlib import Path

from filex.exceptions import FileXError


def find_libreoffice() -> Path:
    """Find the LibreOffice executable."""
    executable = shutil.which("soffice")
    if executable:
        return Path(executable)

    candidates = [
        Path("C:/Program Files/LibreOffice/program/soffice.exe"),
        Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            return candidate

    raise FileXError(
        "LibreOffice was not found. Install LibreOffice and make sure "
        "soffice is available in PATH, or install it in a standard location."
    )


class OfficeConverter:
    """Convert supported document formats using local conversion engines."""

    def convert(
        self,
        input_path: Path,
        output_path: Path,
        target_format: str,
        mode: str = "editable",
    ) -> None:
        input_path = input_path.resolve()
        output_path = output_path.resolve()
        source_format = input_path.suffix.lower().lstrip(".")

        if source_format == "pdf" and target_format == "docx":
            if mode == "visual":
                self._pdf_to_docx_visual(input_path, output_path)
            elif mode == "editable":
                self._pdf_to_docx(input_path, output_path)
            else:
                raise FileXError(f"Unsupported PDF-to-DOCX mode: {mode}")
            return

        self._convert_with_libreoffice(
            input_path=input_path,
            output_path=output_path,
            target_format=target_format,
        )

    def _convert_with_libreoffice(
        self,
        input_path: Path,
        output_path: Path,
        target_format: str,
    ) -> None:
        """Convert a document to any target format supported by LibreOffice."""
        soffice = find_libreoffice()

        with tempfile.TemporaryDirectory(prefix="filex_") as temp:
            temp_dir = Path(temp)
            result_dir = temp_dir / "result"
            profile_dir = temp_dir / "profile"
            result_dir.mkdir()
            profile_dir.mkdir()

            command = [
                str(soffice),
                "--headless",
                f"-env:UserInstallation={profile_dir.as_uri()}",
                "--convert-to",
                target_format,
                "--outdir",
                str(result_dir),
                str(input_path),
            ]

            try:
                process = subprocess.run(
                    command,
                    capture_output=True,
                    text=True,
                    timeout=120,
                    check=False,
                )
            except subprocess.TimeoutExpired as exc:
                raise FileXError(
                    "Document conversion timed out after 120 seconds."
                ) from exc
            except OSError as exc:
                raise FileXError(
                    f"Could not start LibreOffice: {exc}"
                ) from exc

            generated_file = result_dir / f"{input_path.stem}.{target_format}"
            if process.returncode != 0 or not generated_file.is_file():
                details = (process.stderr or process.stdout).strip()
                raise FileXError(
                    f"Conversion from {input_path.suffix.lstrip('.').upper()} "
                    f"to {target_format.upper()} failed. LibreOffice may not "
                    "support this format pair."
                    + (f" Details: {details}" if details else "")
                )

            self._save_without_overwrite(generated_file, output_path)

    def _pdf_to_docx(
        self,
        input_path: Path,
        output_path: Path,
    ) -> None:
        """Convert PDF to editable DOCX using pdf2docx."""
        from pdf2docx import Converter

        with tempfile.TemporaryDirectory(prefix="filex_") as temp:
            generated_file = Path(temp) / "converted.docx"
            converter = None
            try:
                converter = Converter(str(input_path))
                converter.convert(str(generated_file))
            except Exception as exc:
                raise FileXError(
                    f"PDF-to-DOCX conversion failed: {exc}"
                ) from exc
            finally:
                if converter is not None:
                    converter.close()

            if not generated_file.is_file():
                raise FileXError(
                    "PDF-to-DOCX conversion did not produce an output file."
                )
            self._save_without_overwrite(generated_file, output_path)

    def _pdf_to_docx_visual(
        self,
        input_path: Path,
        output_path: Path,
    ) -> None:
        """Render PDF pages as images inside a DOCX document."""
        try:
            import pymupdf
            from docx import Document
            from docx.shared import Inches

            with tempfile.TemporaryDirectory(prefix="filex_visual_") as temp:
                temp_dir = Path(temp)
                generated_file = temp_dir / "converted.docx"
                document = Document()

                with pymupdf.open(input_path) as pdf_document:
                    if pdf_document.page_count == 0:
                        raise FileXError("The input PDF contains no pages.")

                    for page in pdf_document:
                        pixmap = page.get_pixmap(
                            matrix=pymupdf.Matrix(2, 2),
                            alpha=False,
                        )
                        image_path = temp_dir / f"page_{page.number}.png"
                        pixmap.save(image_path)
                        document.add_picture(
                            str(image_path),
                            width=Inches(6.5),
                        )

                document.save(generated_file)
                self._save_without_overwrite(generated_file, output_path)
        except FileXError:
            raise
        except Exception as exc:
            raise FileXError(
                f"PDF-to-DOCX visual conversion failed: {exc}"
            ) from exc

    @staticmethod
    def _save_without_overwrite(
        source_path: Path,
        output_path: Path,
    ) -> None:
        """Copy a converted file without overwriting an existing file."""
        output_created = False
        try:
            with source_path.open("rb") as source:
                with output_path.open("xb") as destination:
                    output_created = True
                    shutil.copyfileobj(source, destination)
        except FileExistsError as exc:
            raise FileXError(
                f"Output file already exists: {output_path}. "
                "Choose another output path."
            ) from exc
        except OSError as exc:
            if output_created:
                output_path.unlink(missing_ok=True)
            raise FileXError(
                f"Could not save converted document: {exc}"
            ) from exc
